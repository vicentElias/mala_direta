import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import psycopg
from app.config import (
    DB_HOST,
    DB_PORT,
    DB_NAME,
    DB_USER,
    DB_PASSWORD,
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

ARQUIVO_EXCEL = Path("C:\\Projetos\\mala_direta\\data\\Matriz de Contato Atualizada 09-2022.xlsx")
ARQUIVO_SAIDA = Path("C:\\Projetos\\mala_direta\\data\\contatos_para_conferencia.xlsx")

# ============================================================
# CAIXAS GENÉRICAS QUE NÃO DEVEM SER TRATADAS COMO PESSOAS
# ============================================================

PALAVRAS_GENERICAS = {
    "contato", "contact", "contatos",
    "financeiro", "financeira", "finance",
    "faturamento", "fiscal",
    "vendas", "venda", "comercial",
    "compras", "compra",
    "rh", "recursos", "administrativo",
    "administracao", "admin",
    "suporte", "support", "atendimento",
    "sac", "marketing", "info",
    "informacao", "informacoes", "geral",
    "diretoria", "recepcao", "ouvidoria",
    "pedidos", "pedido", "logistica",
    "estoque", "expedicao", "nfe", "nf",
    "nota", "notas", "contasapagar",
    "contasareceber", "danfe",
}


# ============================================================
# FUNÇÕES DE LIMPEZA
# ============================================================

def limpar_email(email):
    if email is None:
        return ""

    email = str(email).strip()
    email = email.strip("<>;,.")

    # Remove espaços internos
    email = re.sub(r"\s+", "", email)

    return email.lower()


def email_valido(email):
    padrao = r"^[^@\s]+@[^@\s]+\.[^@\s]+$"
    return bool(re.match(padrao, email))


def remover_acentos(texto):
    texto = unicodedata.normalize("NFD", texto)

    return "".join(
        caractere
        for caractere in texto
        if unicodedata.category(caractere) != "Mn"
    )


def extrair_nome(email):
    """
    Extrai um nome provável da parte anterior ao @.
    Retorna None para caixas genéricas ou identificadores
    que não parecem representar nomes.
    """

    if not email or "@" not in email:
        return None

    parte_local = email.split("@", 1)[0]
    texto = remover_acentos(parte_local).lower()

    # Identificadores com números ou separadores de função
    # normalmente representam setores, não pessoas.
    if re.search(r"\d", texto):
        return None

    if any(separador in texto for separador in [".", "_", "-"]):
        palavras = re.split(r"[._-]+", texto)
    else:
        palavras = [texto]

    palavras = [
        re.sub(r"[^a-z]", "", palavra)
        for palavra in palavras
    ]

    palavras = [palavra for palavra in palavras if palavra]

    if not palavras:
        return None

    if any(palavra in PALAVRAS_GENERICAS for palavra in palavras):
        return None

    # Uma única palavra pode ser nome, mas não temos
    # como confirmar isso apenas pelo endereço.
    if len(palavras) == 1:
        palavra = palavras[0]

        if len(palavra) < 3:
            return None

        return palavra.capitalize()

    return " ".join(palavra.capitalize() for palavra in palavras)


# ============================================================
# LOCALIZAR COLUNAS
# ============================================================

def encontrar_coluna(cabecalhos, nome_procurado):
    procurado = nome_procurado.strip().lower()

    for indice, cabecalho in enumerate(cabecalhos):
        if cabecalho is None:
            continue

        if str(cabecalho).strip().lower() == procurado:
            return indice

    return None


# ============================================================
# LER E CONSOLIDAR OS CONTATOS
# ============================================================

def ler_contatos():
    if not ARQUIVO_EXCEL.exists():
        raise FileNotFoundError(
            f"Arquivo não encontrado: {ARQUIVO_EXCEL.resolve()}"
        )

    workbook = openpyxl.load_workbook(
        ARQUIVO_EXCEL,
        read_only=True,
        data_only=True,
    )

    contatos = {}
    total_ocorrencias = 0
    total_invalidos = 0
    duplicados = 0

    try:
        for nome_aba in workbook.sheetnames:
            planilha = workbook[nome_aba]
            linhas = planilha.iter_rows(values_only=True)

            try:
                cabecalhos = next(linhas)
            except StopIteration:
                continue

            coluna_empresa = encontrar_coluna(
                cabecalhos, "grupo pessoa"
            )

            coluna_email = encontrar_coluna(
                cabecalhos, "e-mail"
            )

            if coluna_empresa is None or coluna_email is None:
                print(
                    f"Aviso: colunas esperadas não encontradas "
                    f"na aba {nome_aba}."
                )
                continue

            for linha in linhas:
                empresa = linha[coluna_empresa]
                valor_email = linha[coluna_email]

                empresa = (
                    str(empresa).strip()
                    if empresa is not None
                    else ""
                )

                if valor_email is None:
                    continue

                for email_original in str(valor_email).split(";"):
                    if not email_original.strip():
                        continue

                    total_ocorrencias += 1
                    email = limpar_email(email_original)

                    if not email_valido(email):
                        total_invalidos += 1
                        continue

                    if email in contatos:
                        duplicados += 1
                        continue

                    nome = extrair_nome(email)

                    contatos[email] = {
                        "nome": nome or "",
                        "empresa": empresa,
                        "email": email,
                        "tipo": "PESSOA" if nome else "GENÉRICO / REVISAR",
                    }

    finally:
        workbook.close()

    return (
        contatos,
        total_ocorrencias,
        duplicados,
        total_invalidos,
    )


# ============================================================
# GERAR O EXCEL DE CONFERÊNCIA
# ============================================================

def gerar_excel():
    (
        contatos,
        total_ocorrencias,
        duplicados,
        invalidos,
    ) = ler_contatos()

    ARQUIVO_SAIDA.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = openpyxl.Workbook()
    planilha = workbook.active
    planilha.title = "Contatos"

    cabecalhos = [
        "Nome",
        "Empresa",
        "E-mail",
        "Tipo de contato",
    ]

    planilha.append(cabecalhos)

    for contato in sorted(
        contatos.values(),
        key=lambda item: (
            item["empresa"].lower(),
            item["email"].lower(),
        ),
    ):
        planilha.append([
            contato["nome"],
            contato["empresa"],
            contato["email"],
            contato["tipo"],
        ])

    # Cabeçalho formatado
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        celula.alignment = Alignment(
            horizontal="center"
        )

    # Largura das colunas
    larguras = {
        "A": 28,
        "B": 45,
        "C": 42,
        "D": 24,
    }

    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions

    workbook.save(ARQUIVO_SAIDA)

    pessoas = sum(
        1 for item in contatos.values()
        if item["tipo"] == "PESSOA"
    )

    genericos = len(contatos) - pessoas

    print()
    print("=" * 60)
    print("RELATÓRIO DE CONFERÊNCIA GERADO")
    print("=" * 60)
    print(f"Ocorrências de e-mail: {total_ocorrencias}")
    print(f"E-mails únicos:        {len(contatos)}")
    print(f"Duplicados:            {duplicados}")
    print(f"E-mails inválidos:     {invalidos}")
    print(f"Prováveis pessoas:     {pessoas}")
    print(f"Genéricos/revisar:     {genericos}")
    print()
    print(f"Arquivo salvo em: {ARQUIVO_SAIDA.resolve()}")
    print()
    print("Nenhum dado foi gravado no PostgreSQL.")
    print("Nenhum e-mail foi enviado.")


if __name__ == "__main__":
    gerar_excel()